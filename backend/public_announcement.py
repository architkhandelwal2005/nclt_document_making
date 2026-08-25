"""Local-first Public Announcement workflow and Form A document generation.

The generated Form A and the newspaper's published copy are deliberately
different records.  Extraction is conservative and always feeds a human review
screen; it never updates operational case data by itself.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, Optional
import io
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader


FORM_FIELDS = (
    "corporate_debtor_name", "date_of_incorporation", "registration_authority", "cin",
    "registered_and_principal_address", "cirp_commencement_date", "order_upload_date",
    "estimated_closure_date", "irp_name", "irp_registration_number",
    "irp_registered_address", "irp_registered_email", "correspondence_address",
    "process_specific_email", "claims_submission_last_date", "creditor_classes",
    "authorised_representatives", "forms_weblink", "authorised_representative_details",
    "announcement_date", "announcement_place", "afa_valid_until",
    "newspaper_name", "publication_language", "edition_or_place", "publication_date",
)


def _display_date(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        return date.fromisoformat(raw[:10]).strftime("%d-%m-%Y")
    except ValueError:
        return raw


def _iso_date(value: str) -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _add_days(value: Any, days: int) -> str:
    raw = str(value or "")[:10]
    try:
        return (date.fromisoformat(raw) + timedelta(days=days)).isoformat()
    except ValueError:
        return ""


def _first(mapping: Dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = mapping.get(key)
        if value not in (None, ""):
            return value
    return ""


def clean_registered_address(value: Any) -> str:
    """Remove extraction spill-over after the registered-office address.

    Older admission-order extraction could append the opening words of the
    following share-capital sentence to the address.  This boundary is not
    part of the registered office and must never reach Form A row 5.
    """
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    text = re.split(
        r"\.\s+The\s+(?:Nominal|Authori[sz]ed|Paid[- ]?Up)\s+Share\s+Capital\b",
        text,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    return text.rstrip(" ,.;:-")


def build_form_defaults(case: Dict[str, Any], profile: Dict[str, Any], intake: Optional[Dict[str, Any]]) -> tuple[Dict[str, Any], Dict[str, Any]]:
    """Build editable defaults while retaining a field-level source map."""
    review = (intake or {}).get("review") or {}
    irp = review.get("irp") or {}
    mca = review.get("mca") or case.get("values", {}).get("mca_master") or {}
    commencement = _first(case, "commencement_date", "order_date")
    order_upload = _first(case.get("values", {}), "order_upload_date")
    claim_trigger = order_upload or commencement
    values = {
        "corporate_debtor_name": case.get("name", ""),
        "date_of_incorporation": _first(mca, "date_of_incorporation", "incorporation_date"),
        "registration_authority": _first(mca, "registrar_of_companies", "roc", "roc_location", "registration_authority"),
        "cin": case.get("cin", ""),
        "registered_and_principal_address": clean_registered_address(case.get("registered_address", "")),
        "cirp_commencement_date": commencement,
        "order_upload_date": order_upload,
        "estimated_closure_date": _add_days(commencement, 180),
        "irp_name": _first(profile, "ip_name") or _first(irp, "name"),
        "irp_registration_number": _first(profile, "ibbi_reg_no") or _first(irp, "registration_number"),
        "irp_registered_address": _first(profile, "ip_reg_address") or _first(irp, "address"),
        "irp_registered_email": _first(profile, "ip_email") or _first(irp, "email"),
        "correspondence_address": _first(profile, "ip_reg_address") or _first(irp, "address"),
        "process_specific_email": _first(profile, "process_email"),
        "claims_submission_last_date": _add_days(claim_trigger, 14),
        "creditor_classes": "Based on information available, no class of creditors under section 21(6A)(b) has been ascertained.",
        "authorised_representatives": "Not applicable.",
        "forms_weblink": "https://ibbi.gov.in/en/home/downloads",
        "authorised_representative_details": "Not applicable based on information presently available with the IRP.",
        "announcement_date": date.today().isoformat(),
        "announcement_place": "",
        "afa_valid_until": _first(profile, "afa_validity") or _first(irp, "afa_valid_until"),
        "newspaper_name": "", "publication_language": "", "edition_or_place": "", "publication_date": "",
    }
    provenance = {}
    for field in values:
        if field in {"corporate_debtor_name", "cin", "registered_and_principal_address", "cirp_commencement_date"}:
            provenance[field] = {"source_type": "admission_order", "source_id": (intake or {}).get("document_id") or (intake or {}).get("id")}
        elif field in {"date_of_incorporation", "registration_authority"}:
            provenance[field] = {"source_type": "company_master", "source_id": case.get("id")}
        elif field.startswith("irp_") or field in {"process_specific_email", "correspondence_address", "afa_valid_until"}:
            provenance[field] = {"source_type": "professional_master", "source_id": profile.get("user_id", "")}
        elif field in {"estimated_closure_date", "claims_submission_last_date"}:
            provenance[field] = {"source_type": "deterministic_calculation", "source_id": "commencement/order-upload-date"}
        else:
            provenance[field] = {"source_type": "user_review", "source_id": ""}
    return values, provenance


def _set_cell_margins(cell, top=60, start=80, bottom=60, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(amount))
        tag.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def _font(run, size=10.5, bold=False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def _paragraph(container, text="", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10.5, bold=False, before=0, after=3):
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(str(text or ""))
    _font(run, size, bold)
    return paragraph


FORM_ROWS = (
    ("Name of corporate debtor", "corporate_debtor_name"),
    ("Date of incorporation of corporate debtor", "date_of_incorporation"),
    ("Authority under which corporate debtor is incorporated / registered", "registration_authority"),
    ("Corporate Identity Number of corporate debtor", "cin"),
    ("Address of the registered office and principal office (if any) of corporate debtor", "registered_and_principal_address"),
    ("Insolvency commencement date in respect of corporate debtor", "cirp_commencement_date"),
    ("Estimated date of closure of insolvency resolution process", "estimated_closure_date"),
    ("Name and registration number of the insolvency professional acting as interim resolution professional", "irp_identity"),
    ("Address and e-mail of the interim resolution professional, as registered with the Board", "irp_registered_contact"),
    ("Address and e-mail to be used for correspondence with the interim resolution professional", "correspondence_contact"),
    ("Last date for submission of claims", "claims_submission_last_date"),
    ("Classes of creditors, if any, under clause (b) of sub-section (6A) of section 21, ascertained by the interim resolution professional", "creditor_classes"),
    ("Names of Insolvency Professionals identified to act as Authorised Representative of creditors in a class (three names for each class)", "authorised_representatives"),
    ("(a) Relevant Forms and (b) Details of authorised representatives are available at:", "forms_and_representatives"),
)


def generate_form_a(values: Dict[str, Any], output: Path) -> Path:
    """Create an editable legal-size DOCX following the clean two-page Form A."""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(14)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)

    _paragraph(doc, "FORM A", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, after=0)
    _paragraph(doc, "PUBLIC ANNOUNCEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, after=0)
    _paragraph(doc, "(Under Regulation 6 of the Insolvency and Bankruptcy Board of India (Insolvency Resolution Process for Corporate Persons) Regulations, 2016)", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, after=8)
    attention = _paragraph(doc, "FOR THE ATTENTION OF THE CREDITORS OF ", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, after=2)
    run = attention.add_run(str(values.get("corporate_debtor_name") or "").upper())
    _font(run, 11, True)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = (Inches(0.45), Inches(3.55), Inches(3.05))
    header = table.rows[0]
    merged = header.cells[0].merge(header.cells[2])
    merged.text = "RELEVANT PARTICULARS"
    merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(merged.paragraphs[0].runs[0], 10.5, True)
    _set_repeat_header(header)

    computed = dict(values)
    computed["registered_and_principal_address"] = clean_registered_address(
        values.get("registered_and_principal_address")
    )
    computed["irp_identity"] = f"{values.get('irp_name','')}\nReg. No.: {values.get('irp_registration_number','')}".strip()
    computed["irp_registered_contact"] = f"Address: {values.get('irp_registered_address','')}\nEmail: {values.get('irp_registered_email','')}".strip()
    computed["correspondence_contact"] = f"Address: {values.get('correspondence_address','')}\nEmail: {values.get('process_specific_email','')}".strip()
    computed["forms_and_representatives"] = f"(a) Weblink: {values.get('forms_weblink','')}\nPhysical address: As stated at item 10.\n(b) {values.get('authorised_representative_details','')}".strip()
    date_keys = {"date_of_incorporation", "cirp_commencement_date", "estimated_closure_date", "claims_submission_last_date"}
    for index, (label, key) in enumerate(FORM_ROWS, 1):
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell)
        cells[0].text = f"{index}."
        cells[1].text = label
        value = computed.get(key, "")
        cells[2].text = _display_date(value) if key in date_keys else str(value or "")
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    _font(run, 9.4)

    commencement = _display_date(values.get("cirp_commencement_date"))
    upload_note = f" (order uploaded on {_display_date(values.get('order_upload_date'))})" if values.get("order_upload_date") else ""
    _paragraph(doc, f"Notice is hereby given that the National Company Law Tribunal has ordered the commencement of a corporate insolvency resolution process of {values.get('corporate_debtor_name','')} on {commencement}{upload_note}.", before=6, after=6)
    doc.add_page_break()
    deadline = _display_date(values.get("claims_submission_last_date"))
    _paragraph(doc, f"The creditors of {values.get('corporate_debtor_name','')} are hereby called upon to submit their claims with proof on or before {deadline} to the interim resolution professional at the address mentioned against entry No. 10.", after=7)
    _paragraph(doc, "The financial creditors shall submit their claims with proof by electronic means only. All other creditors may submit their claims with proof in person, by post or by electronic means.", after=7)
    _paragraph(doc, "A financial creditor belonging to a class, as listed against entry No. 12, shall indicate its choice of authorised representative from among the three insolvency professionals listed against entry No. 13 to act as authorised representative of the class in Form CA.", after=7)
    _paragraph(doc, "Submission of false or misleading proofs of claim shall attract penalties.", after=12)
    signature = _paragraph(doc, values.get("irp_name", ""), align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, after=0)
    _paragraph(doc, "Interim Resolution Professional", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    _paragraph(doc, values.get("irp_registration_number", ""), align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    if values.get("afa_valid_until"):
        _paragraph(doc, f"AFA valid till {_display_date(values.get('afa_valid_until'))}", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    _paragraph(doc, f"Date: {_display_date(values.get('announcement_date'))}", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    _paragraph(doc, f"Place: {values.get('announcement_place','')}", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)

    # Remove the empty default footer paragraph's visible spacing.
    for paragraph in section.footer.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
    doc.save(output)
    return output


def _regex(text: str, patterns: Iterable[str]) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, re.I | re.S)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip(" :-\n\t")
    return ""


def _tesseract_executable() -> Optional[str]:
    executable = shutil.which("tesseract")
    if executable:
        return executable
    conventional = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    return str(conventional) if conventional.exists() else None


def _ocr_pdf_images(reader: PdfReader) -> tuple[str, list[str]]:
    executable = _tesseract_executable()
    if not executable:
        return "", ["The PDF appears scanned, but local Tesseract OCR is not installed. Enter and review the publication fields manually."]
    chunks, warnings = [], []
    with tempfile.TemporaryDirectory(prefix="casefile-pa-ocr-") as temporary:
        temp = Path(temporary)
        for page_number, page in enumerate(reader.pages, 1):
            images = list(page.images)
            if not images:
                warnings.append(f"Page {page_number} contains no directly extractable image for OCR.")
                continue
            # The largest embedded image normally represents the scanned page.
            image = max(images, key=lambda item: len(item.data))
            suffix = Path(image.name or "page.png").suffix or ".png"
            source = temp / f"page-{page_number}{suffix}"
            source.write_bytes(image.data)
            result = subprocess.run([executable, str(source), "stdout", "-l", "eng"], capture_output=True, text=True, timeout=120, check=False)
            if result.returncode == 0:
                chunks.append(result.stdout)
            else:
                warnings.append(f"OCR failed on page {page_number}; manual review is required.")
    return "\n".join(chunks), warnings


def extract_published_pdf(path: Path) -> Dict[str, Any]:
    """Extract possible fields without committing any value to the case."""
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("Password-protected published announcements are not supported")
    direct = "\n".join(page.extract_text() or "" for page in reader.pages)
    warnings: list[str] = []
    method = "pdf_text"
    text = direct
    if len(re.sub(r"\s+", "", direct)) < 120:
        text, warnings = _ocr_pdf_images(reader)
        method = "tesseract_ocr" if text else "manual_required"
    fields = {
        "publication_date": _iso_date(_regex(text, [r"publication\s+date\s*[:\-]?\s*(\d{1,2}[./-]\d{1,2}[./-]\d{4})", r"date\s*[:\-]\s*(\d{1,2}[./-]\d{1,2}[./-]\d{4})"])) or "",
        "newspaper_name": "",
        "publication_language": "",
        "edition_or_place": _regex(text, [r"(?:edition|place)\s*[:\-]\s*([^\n]{2,60})"]),
        "corporate_debtor_name": _regex(text, [r"attention\s+of\s+the\s+creditors\s+of\s+(.+?)\s+relevant\s+particulars", r"name\s+of\s+corporate\s+debtor\s+(.+?)(?:\n\s*2\.|date\s+of\s+incorporation)"]),
        "cin": _regex(text, [r"corporate\s+identity[^\n]{0,80}(?:\n[^\n]{0,80})?\s+(U[A-Z0-9]{15,25})", r"\bCIN\s*[:\-]?\s*([A-Z0-9]{15,25})"]),
        "cirp_commencement_date": _iso_date(_regex(text, [r"(?:insolvency|CIRP)\s+commencement\s+date[^\d]{0,40}(\d{1,2}[./-]\d{1,2}[./-]\d{4})"])) or "",
        "claims_submission_last_date": _iso_date(_regex(text, [r"last\s+date\s+for\s+submission\s+of\s+claims[^\d]{0,40}(\d{1,2}[./-]\d{1,2}[./-]\d{4})", r"on\s+or\s+before\s+(\d{1,2}[./-]\d{1,2}[./-]\d{4})"])) or "",
        "irp_name": _regex(text, [r"(?:interim\s+resolution\s+professional|IRP)\s*[:\-]?\s*(?:Mr\.?|Ms\.?)?\s*([A-Z][A-Za-z .]{3,80})"]),
        "irp_registration_number": re.sub(r"\s+", "", _regex(text, [r"(IBBI/IPA-[A-Z0-9]+/IP-P[A-Z0-9]+/\d{4}-\s*\d{2}/\d+)", r"(IBBI/[A-Z0-9/\-]{10,})"])),
        "process_specific_email": _regex(text, [r"correspondence\s+with\s+the\s+interim\s+resolution\s+professional.+?Email\s*[:\-]?\s*([A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,})", r"Email\s*[:\-]?\s*([A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,})"]),
    }
    return {"method": method, "page_count": len(reader.pages), "text_length": len(text), "fields": fields, "warnings": warnings}
