"""Local, page-aware preparation for claim-bundle documents.

No document bytes are sent by this module.  It extracts text locally and marks
pages needing OCR instead of pretending an unreadable scan contains no facts.
"""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

from docx import Document
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".csv"}


@dataclass(frozen=True)
class PreparedPage:
    number: int
    text: str
    extraction_status: str


@dataclass(frozen=True)
class PreparedDocument:
    document_id: str
    name: str
    path: Path
    sha256: str
    pages: tuple[PreparedPage, ...]

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def text(self) -> str:
        return "\n\n".join(
            f"--- DOCUMENT {self.document_id} | {self.name} | PAGE {page.number} | {page.extraction_status} ---\n{page.text}"
            for page in self.pages
        )


class ClaimDocumentError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _tesseract_executable() -> str | None:
    executable = shutil.which("tesseract")
    if executable:
        return executable
    conventional = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    return str(conventional) if conventional.exists() else None


def _ocr_page(page: Any, page_number: int) -> str:
    executable = _tesseract_executable()
    if not executable:
        return ""
    try:
        images = list(page.images)
        if not images:
            return ""
        image = max(images, key=lambda item: len(item.data))
        with tempfile.TemporaryDirectory(prefix="casefile-claim-ocr-") as temporary:
            suffix = Path(image.name or "page.png").suffix or ".png"
            source = Path(temporary) / f"page-{page_number}{suffix}"
            source.write_bytes(image.data)
            result = subprocess.run([executable, str(source), "stdout", "-l", "eng"], capture_output=True,
                                    text=True, timeout=120, check=False)
            return result.stdout.strip() if result.returncode == 0 else ""
    except Exception:
        return ""


def _pdf_pages(path: Path) -> tuple[PreparedPage, ...]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ClaimDocumentError("DOCUMENT_EXTRACTION_FAILED", "The claim PDF could not be opened") from exc
    pages: list[PreparedPage] = []
    for index, page in enumerate(reader.pages, 1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        status = "TEXT_EXTRACTED" if len(text) >= 40 else "OCR_REQUIRED"
        if status == "OCR_REQUIRED":
            ocr_text = _ocr_page(page, index)
            if len(ocr_text) >= 20:
                text, status = ocr_text, "OCR_EXTRACTED"
        pages.append(PreparedPage(index, text, status))
    if not pages:
        raise ClaimDocumentError("DOCUMENT_EXTRACTION_FAILED", "The claim PDF has no pages")
    return tuple(pages)


def _docx_pages(path: Path) -> tuple[PreparedPage, ...]:
    try:
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    except Exception as exc:
        raise ClaimDocumentError("DOCUMENT_EXTRACTION_FAILED", "The claim DOCX could not be read") from exc
    return (PreparedPage(1, text.strip(), "TEXT_EXTRACTED" if text.strip() else "OCR_REQUIRED"),)


def prepare_document(document: dict[str, Any], data_dir: Path) -> PreparedDocument:
    path = (Path(data_dir).resolve() / str(document.get("storage_path") or "")).resolve()
    if not path.is_file() or Path(data_dir).resolve() not in path.parents:
        raise ClaimDocumentError("DOCUMENT_UNAVAILABLE", "A selected claim document is unavailable")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ClaimDocumentError("UNSUPPORTED_DOCUMENT_TYPE", f"Claim AI cannot read {suffix or 'this file type'}")
    if suffix == ".pdf":
        pages = _pdf_pages(path)
    elif suffix == ".docx":
        pages = _docx_pages(path)
    else:
        try:
            text = path.read_text(encoding="utf-8", errors="replace").strip()
        except OSError as exc:
            raise ClaimDocumentError("DOCUMENT_EXTRACTION_FAILED", "The claim text file could not be read") from exc
        pages = (PreparedPage(1, text, "TEXT_EXTRACTED" if text else "OCR_REQUIRED"),)
    return PreparedDocument(
        document_id=str(document["id"]), name=str(document.get("name") or path.name), path=path,
        sha256=sha256(path.read_bytes()).hexdigest(), pages=pages,
    )


def pages_text(document: PreparedDocument, start_page: int, end_page: int) -> str:
    selected = [page for page in document.pages if start_page <= page.number <= end_page]
    return "\n\n".join(
        f"--- DOCUMENT {document.document_id} | {document.name} | PAGE {page.number} | {page.extraction_status} ---\n{page.text}"
        for page in selected
    )


def page_batches(document: PreparedDocument, target_characters: int = 18_000) -> list[tuple[int, int, str]]:
    """Create controlled, page-boundary batches for free-tier providers."""
    batches: list[tuple[int, int, str]] = []
    current: list[PreparedPage] = []
    size = 0
    for page in document.pages:
        page_size = len(page.text) + 120
        if current and size + page_size > target_characters:
            batches.append((current[0].number, current[-1].number,
                            pages_text(document, current[0].number, current[-1].number)))
            current, size = [], 0
        current.append(page)
        size += page_size
    if current:
        batches.append((current[0].number, current[-1].number,
                        pages_text(document, current[0].number, current[-1].number)))
    return batches
