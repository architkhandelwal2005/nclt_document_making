from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


ROOT = Path(r"E:\nclt document")
REFERENCE_DIR = ROOT / "reference documents"
OUTPUT_DIR = ROOT / "tmp" / "coc_reference_audit"


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")


def points(value):
    return None if value is None else round(value.pt, 2)


def inches(value):
    return None if value is None else round(value.inches, 3)


def safe(callable_value):
    try:
        return callable_value()
    except (AttributeError, TypeError, ValueError):
        return None


def color_hex(run):
    value = run.font.color.rgb
    return None if value is None else str(value)


def run_record(run):
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "font": run.font.name,
        "size_pt": points(run.font.size),
        "color": color_hex(run),
        "all_caps": run.font.all_caps,
        "small_caps": run.font.small_caps,
        "superscript": run.font.superscript,
        "subscript": run.font.subscript,
    }


def paragraph_record(paragraph):
    fmt = paragraph.paragraph_format
    numbering = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    return {
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": None if paragraph.alignment is None else str(paragraph.alignment),
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
        "space_before_pt": points(fmt.space_before),
        "space_after_pt": points(fmt.space_after),
        "line_spacing": str(fmt.line_spacing) if fmt.line_spacing is not None else None,
        "left_indent_in": inches(safe(lambda: fmt.left_indent)),
        "right_indent_in": inches(safe(lambda: fmt.right_indent)),
        "first_line_indent_in": inches(safe(lambda: fmt.first_line_indent)),
        "numbered": numbering is not None,
        "runs": [run_record(run) for run in paragraph.runs],
    }


def cell_record(cell):
    tc_pr = cell._tc.tcPr
    fill = None
    shading = tc_pr.find(qn("w:shd"))
    if shading is not None:
        fill = shading.get(qn("w:fill"))
    width = tc_pr.tcW
    return {
        "text": cell.text,
        "width": None if width is None else width.get(qn("w:w")),
        "width_type": None if width is None else width.get(qn("w:type")),
        "fill": fill,
        "vertical_alignment": None if cell.vertical_alignment is None else str(cell.vertical_alignment),
        "paragraphs": [paragraph_record(p) for p in cell.paragraphs],
    }


def table_record(table, index):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    border_record = {}
    if borders is not None:
        for child in borders:
            border_record[child.tag.split("}")[-1]] = dict(child.attrib)
    grid = table._tbl.tblGrid
    grid_widths = [] if grid is None else [col.get(qn("w:w")) for col in grid]
    return {
        "index": index,
        "style": table.style.name if table.style else None,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "autofit": table.autofit,
        "grid_widths": grid_widths,
        "borders": border_record,
        "row_records": [
            {
                "height": None if row.height is None else row.height.twips,
                "height_rule": None if row.height_rule is None else str(row.height_rule),
                "repeat_header": row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None,
                "cells": [cell_record(cell) for cell in row.cells],
            }
            for row in table.rows
        ],
    }


def container_record(container):
    return {
        "paragraphs": [paragraph_record(p) for p in container.paragraphs],
        "tables": [table_record(t, i) for i, t in enumerate(container.tables)],
    }


def section_record(section, index):
    return {
        "index": index,
        "orientation": "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait",
        "page_width_in": inches(section.page_width),
        "page_height_in": inches(section.page_height),
        "top_margin_in": inches(section.top_margin),
        "bottom_margin_in": inches(section.bottom_margin),
        "left_margin_in": inches(section.left_margin),
        "right_margin_in": inches(section.right_margin),
        "header_distance_in": inches(section.header_distance),
        "footer_distance_in": inches(section.footer_distance),
        "gutter_in": inches(section.gutter),
        "different_first_page": section.different_first_page_header_footer,
        "header_linked": section.header.is_linked_to_previous,
        "footer_linked": section.footer.is_linked_to_previous,
        "first_header_linked": section.first_page_header.is_linked_to_previous,
        "first_footer_linked": section.first_page_footer.is_linked_to_previous,
        "even_header_linked": section.even_page_header.is_linked_to_previous,
        "even_footer_linked": section.even_page_footer.is_linked_to_previous,
        "header": container_record(section.header),
        "footer": container_record(section.footer),
        "first_page_header": container_record(section.first_page_header),
        "first_page_footer": container_record(section.first_page_footer),
        "even_page_header": container_record(section.even_page_header),
        "even_page_footer": container_record(section.even_page_footer),
    }


def package_record(path: Path):
    with zipfile.ZipFile(path) as archive:
        names = sorted(archive.namelist())
        media = []
        for name in names:
            if name.startswith("word/media/"):
                data = archive.read(name)
                media.append({
                    "part": name,
                    "bytes": len(data),
                    "sha256": hashlib.sha256(data).hexdigest(),
                })
        xml_flags = Counter()
        for name in names:
            if not name.endswith(".xml"):
                continue
            text = archive.read(name).decode("utf-8", errors="ignore")
            for label, token in {
                "content_controls": "<w:sdt",
                "fields": "<w:fldSimple",
                "field_chars": "<w:fldChar",
                "bookmarks": "<w:bookmarkStart",
                "comments": "<w:commentRangeStart",
                "tracked_insertions": "<w:ins",
                "tracked_deletions": "<w:del",
                "drawings": "<w:drawing",
                "floating_anchors": "<wp:anchor",
                "inline_images": "<wp:inline",
                "section_properties": "<w:sectPr",
                "manual_page_breaks": 'w:type="page"',
            }.items():
                xml_flags[label] += text.count(token)
        return {"parts": names, "media": media, "xml_feature_counts": dict(xml_flags)}


def inspect(path: Path):
    document = Document(path)
    style_usage = Counter(p.style.name if p.style else "(none)" for p in document.paragraphs)
    core = document.core_properties
    return {
        "source": str(path),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "core_properties": {
            "title": core.title,
            "subject": core.subject,
            "author": core.author,
            "last_modified_by": core.last_modified_by,
            "created": core.created.isoformat() if core.created else None,
            "modified": core.modified.isoformat() if core.modified else None,
            "revision": core.revision,
        },
        "section_count": len(document.sections),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "style_usage": dict(style_usage),
        "sections": [section_record(s, i) for i, s in enumerate(document.sections)],
        "paragraphs": [paragraph_record(p) for p in document.paragraphs],
        "tables": [table_record(t, i) for i, t in enumerate(document.tables)],
        "package": package_record(path),
    }


def main():
    for path in sorted(REFERENCE_DIR.glob("*.docx")):
        target_dir = OUTPUT_DIR / safe_name(path.stem)
        target_dir.mkdir(parents=True, exist_ok=True)
        output = target_dir / "structure.json"
        output.write_text(json.dumps(inspect(path), indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"{path.name} -> {output}")


if __name__ == "__main__":
    main()
